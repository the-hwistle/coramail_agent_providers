from __future__ import annotations

import mimetypes
import re
from pathlib import Path
from typing import Any, Literal
from uuid import UUID

from app.document_processing.attachment_classifier import (
    STRUCTURED_DOCUMENT_TYPES,
    canonical_document_category,
    document_category_label,
    infer_attachment_document_category,
)
from app.document_processing.paddleocr_adapter import paddleocr_mode, parse_with_paddleocr
from app.schemas.attachment_analysis import (
    AttachmentAnalysisResult,
    AttachmentAnalysisStatus,
    AttachmentEvidence,
    AttachmentPage,
    AttachmentTable,
)


class AttachmentParserError(RuntimeError):
    pass


def resolve_storage_path(storage_uri: str, project_dir: Path) -> Path:
    value = storage_uri.strip()
    if value.startswith("file://"):
        return Path(value[7:]).expanduser().resolve()
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = project_dir / path
    return path.resolve()


class AttachmentParserDispatcher:
    def __init__(self, project_dir: Path):
        self.project_dir = project_dir

    def analyze(self, attachment: dict[str, Any]) -> AttachmentAnalysisResult:
        attachment_id = UUID(str(attachment["id"]))
        filename = str(attachment.get("filename") or "attachment")
        content_type = str(attachment.get("content_type") or mimetypes.guess_type(filename)[0] or "application/octet-stream")
        path = resolve_storage_path(str(attachment.get("storage_uri") or ""), self.project_dir)
        if not path.exists() or not path.is_file():
            return AttachmentAnalysisResult(
                attachment_id=attachment_id,
                filename=filename,
                content_type=content_type,
                status=AttachmentAnalysisStatus.FAILED,
                warnings=["attachment_source_missing"],
                error_message=f"attachment file not found: {path}",
            )
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf" or content_type == "application/pdf":
                return self._pdf(attachment_id, filename, content_type, path)
            if suffix in {".xlsx", ".xlsm"}:
                return self._xlsx(attachment_id, filename, content_type, path)
            if suffix == ".docx":
                return self._docx(attachment_id, filename, content_type, path)
            if content_type.startswith("image/") or suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}:
                return self._image(attachment_id, filename, content_type, path)
            if content_type.startswith("text/") or suffix in {".txt", ".csv"}:
                return self._text(attachment_id, filename, content_type, path)
            return AttachmentAnalysisResult(
                attachment_id=attachment_id,
                filename=filename,
                content_type=content_type,
                status=AttachmentAnalysisStatus.UNSUPPORTED,
                warnings=[f"unsupported_attachment_type:{suffix or content_type}"],
            )
        except Exception as exc:
            return AttachmentAnalysisResult(
                attachment_id=attachment_id,
                filename=filename,
                content_type=content_type,
                status=AttachmentAnalysisStatus.FAILED,
                error_message=f"{type(exc).__name__}: {exc}",
            )

    def _pdf(self, attachment_id: UUID, filename: str, content_type: str, path: Path) -> AttachmentAnalysisResult:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages: list[AttachmentPage] = []
        warnings: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            pages.append(AttachmentPage(page_number=index, text=text))
            if not text:
                warnings.append(f"page_{index}_requires_vision_or_ocr")
        text = "\n\n".join(page.text for page in pages if page.text)
        ocr = None
        if paddleocr_mode() == "always" or not text.strip() or warnings:
            ocr = parse_with_paddleocr(path)
            if ocr and ocr.text.strip():
                text = ocr.text
                pages = [
                    AttachmentPage(page_number=index, text=page_text)
                    for index, page_text in enumerate(ocr.page_texts, start=1)
                    if page_text.strip()
                ]
                warnings = [warning for warning in warnings if not warning.endswith("requires_vision_or_ocr")]
                warnings.extend(warning for warning in ocr.warnings if warning not in warnings)
            elif ocr:
                warnings.extend(warning for warning in ocr.warnings if warning not in warnings)
        unresolved_page_warnings = [warning for warning in warnings if warning.endswith("requires_vision_or_ocr")]
        status = AttachmentAnalysisStatus.COMPLETED if text and not unresolved_page_warnings else AttachmentAnalysisStatus.PARTIAL_SUCCESS
        if not text:
            status = AttachmentAnalysisStatus.PARTIAL_SUCCESS
        category, confidence, reason = infer_attachment_document_category(
            filename=filename,
            content_type=content_type,
            extracted_text=text,
        )
        if reason != "no_rule_match":
            warnings.append(reason)
        return self._result(
            attachment_id,
            filename,
            content_type,
            status,
            text,
            pages=pages,
            warnings=warnings,
            document_type=category,
            document_type_confidence=confidence,
            document_html=ocr.html if ocr else "",
        )

    def _xlsx(self, attachment_id: UUID, filename: str, content_type: str, path: Path) -> AttachmentAnalysisResult:
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True, read_only=True)
        tables: list[AttachmentTable] = []
        sections: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[list[Any]] = []
            for row in sheet.iter_rows(values_only=True):
                values = [value for value in row]
                if any(value not in (None, "") for value in values):
                    rows.append(values)
            if rows:
                tables.append(AttachmentTable(name=sheet.title, sheet_name=sheet.title, rows=rows))
                sections.append(f"[{sheet.title}]\n" + "\n".join(" | ".join("" if value is None else str(value) for value in row) for row in rows))
        text = "\n\n".join(sections)
        return self._result(attachment_id, filename, content_type, AttachmentAnalysisStatus.COMPLETED, text, tables=tables)

    def _docx(self, attachment_id: UUID, filename: str, content_type: str, path: Path) -> AttachmentAnalysisResult:
        from docx import Document

        document = Document(path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        tables: list[AttachmentTable] = []
        for index, table in enumerate(document.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            tables.append(AttachmentTable(name=f"table_{index}", rows=rows))
        table_text = ["\n".join(" | ".join(row) for row in table.rows) for table in tables]
        text = "\n\n".join(paragraphs + table_text)
        return self._result(attachment_id, filename, content_type, AttachmentAnalysisStatus.COMPLETED, text, tables=tables)

    def _image(self, attachment_id: UUID, filename: str, content_type: str, path: Path) -> AttachmentAnalysisResult:
        from PIL import Image

        ocr = parse_with_paddleocr(path)
        if ocr and ocr.text.strip():
            return self._result(
                attachment_id,
                filename,
                content_type,
                AttachmentAnalysisStatus.COMPLETED,
                ocr.text,
                pages=[
                    AttachmentPage(page_number=index, text=page_text)
                    for index, page_text in enumerate(ocr.page_texts, start=1)
                    if page_text.strip()
                ],
                warnings=ocr.warnings,
                document_html=ocr.html,
            )

        with Image.open(path) as image:
            fields = {"width": image.width, "height": image.height, "mode": image.mode, "format": image.format}
        category, confidence, reason = infer_attachment_document_category(
            filename=filename,
            content_type=content_type,
        )
        return AttachmentAnalysisResult(
            attachment_id=attachment_id,
            filename=filename,
            content_type=content_type,
            status=AttachmentAnalysisStatus.PARTIAL_SUCCESS,
            document_type=category,
            document_type_confidence=confidence,
            analysis_summary=_summary_for_category(category),
            fields=fields,
            warnings=["vision_analysis_required", reason, *(ocr.warnings if ocr else [])],
        )

    def _text(self, attachment_id: UUID, filename: str, content_type: str, path: Path) -> AttachmentAnalysisResult:
        text = path.read_text(encoding="utf-8", errors="replace")
        return self._result(attachment_id, filename, content_type, AttachmentAnalysisStatus.COMPLETED, text)

    @staticmethod
    def _result(
        attachment_id: UUID,
        filename: str,
        content_type: str,
        status: AttachmentAnalysisStatus,
        text: str,
        *,
        pages: list[AttachmentPage] | None = None,
        tables: list[AttachmentTable] | None = None,
        warnings: list[str] | None = None,
        document_type: str | None = None,
        document_type_confidence: float = 0.0,
        document_html: str = "",
    ) -> AttachmentAnalysisResult:
        result_warnings = list(warnings or [])
        if not document_type:
            category, confidence, reason = infer_attachment_document_category(
                filename=filename,
                content_type=content_type,
                extracted_text=text,
            )
            document_type = category
            document_type_confidence = confidence
            if reason != "no_rule_match":
                result_warnings.append(reason)

        evidence = []
        if text.strip():
            evidence.append(AttachmentEvidence(attachment_id=attachment_id, text=text[:4000]))
        analysis_summary = _summary_for_category(document_type) if document_type else ""
        fields, validation_warnings = validated_predefined_document_fields(document_type, text, {})
        result_warnings.extend(validation_warnings)
        return AttachmentAnalysisResult(
            attachment_id=attachment_id,
            filename=filename,
            content_type=content_type,
            status=status,
            document_type=document_type,
            document_type_confidence=document_type_confidence,
            analysis_summary=analysis_summary,
            extracted_text=text,
            document_html=document_html,
            pages=pages or [],
            tables=tables or [],
            fields=fields,
            evidence=evidence,
            warnings=result_warnings,
        )


def _summary_for_category(document_type: str | None) -> str:
    return f"{document_category_label(document_type or 'unknown')} 유형으로 추정됩니다."


PREDEFINED_DOCUMENT_FIELD_ORDER = {
    "quote": [
        "To", "Attn", "Your Ref No", "Date", "Our Ref No", "In Charge", "Tel", "Vessel",
        "Total Price", "line_items",
    ],
    "rfq": [
        "To", "Attn", "Email", "Fax", "Vessel", "Date", "Our Ref No", "In Charge",
        "Tel", "line_items",
    ],
    "payment_request": [
        "To", "Attn", "Add", "Tel", "Your Ref. No.", "공급자 - 등록번호", "공급자 - 상호",
        "공급자 - 소재지", "공급자 - 계좌번호", "입금요청일", "No.", "Description", "Qty",
        "Unit", "U/Price", "Amount", "소계(VAT제외)", "합계(VAT포함)",
    ],
    "transaction_statement": [
        "VESSEL", "YOUR REF NO", "출고일", "공급받는자-등록번호", "공급받는자-상호(법인명)",
        "공급받는자-사업장 주소", "공급받는자-전화번호", "공급자-등록번호",
        "공급자-상호(법인명)", "공급자-성명(대표자)", "공급자-사업장 주소", "공급자-업태",
        "공급자-종목", "공급자-계좌번호", "합계금액", "No", "Description", "Qty", "Unit",
        "U/Price", "Amount", "공급가액", "세액", "비고", "인수자",
    ],
}

PredefinedDocumentFieldName = Literal[
    "To",
    "Attn",
    "Your Ref No",
    "Vessel",
    "Date",
    "Our Ref No",
    "In Charge",
    "Tel",
    "Total Price",
    "No",
    "Description",
    "Code",
    "Qty",
    "Unit",
    "U/Price",
    "Amount",
    "REMARK",
    "Terms & Conditions - Delivery time",
    "Email",
    "Fax",
    "Add",
    "Your Ref. No.",
    "공급자 - 등록번호",
    "공급자 - 상호",
    "공급자 - 소재지",
    "공급자 - 계좌번호",
    "입금요청일",
    "No.",
    "소계(VAT제외)",
    "합계(VAT포함)",
    "VESSEL",
    "YOUR REF NO",
    "출고일",
    "공급받는자-등록번호",
    "공급받는자-상호(법인명)",
    "공급받는자-사업장 주소",
    "공급받는자-전화번호",
    "공급자-등록번호",
    "공급자-상호(법인명)",
    "공급자-성명(대표자)",
    "공급자-사업장 주소",
    "공급자-업태",
    "공급자-종목",
    "공급자-계좌번호",
    "합계금액",
    "공급가액",
    "세액",
    "비고",
    "인수자",
    "line_items",
]

QUOTE_HEADER_FIELDS = (
    "To",
    "Attn",
    "Your Ref No",
    "Date",
    "Our Ref No",
    "In Charge",
    "Tel",
    "Vessel",
    "Total Price",
)

QUOTE_LINE_ITEM_FIELDS = ("Description", "Qty", "Unit", "U/Price", "Amount")
RFQ_LINE_ITEM_FIELDS = ("No", "Description", "Code", "Qty", "Unit")

FORM_FIELD_LABELS = {
    "quote": [
        "Terms & Conditions - Delivery time",
        "견적 유효기간",
        "견적번호",
        "견적일자",
        "수신 회사",
        "합계금액",
        "Your Ref. No.",
        "Your Ref No",
        "Our Ref. No.",
        "Our Ref No",
        "Reference No.",
        "Total Price",
        "Grand Total",
        "Delivery Time",
        "In Charge",
        "Vessel Name",
        "Vessel",
        "Attn",
        "Date",
        "Email",
        "Fax",
        "Tel",
        "To",
    ],
    "rfq": [
        "Your Ref. No.",
        "Your Ref No",
        "Our Ref. No.",
        "Our Ref No",
        "Reference No.",
        "In Charge",
        "Vessel Name",
        "Vessel",
        "Attn",
        "Date",
        "Email",
        "Fax",
        "Tel",
        "To",
    ],
    "payment_request": [
        "공급자 - 등록번호",
        "공급자 - 계좌번호",
        "공급자 - 소재지",
        "공급자 - 상호",
        "합계(VAT포함)",
        "소계(VAT제외)",
        "입금요청일",
        "Your Ref. No.",
        "Description",
        "Amount",
        "U/Price",
        "Unit",
        "Qty",
        "Attn",
        "Date",
        "Tel",
        "Add",
        "To",
    ],
    "transaction_statement": [
        "공급받는자-등록번호",
        "공급받는자-상호(법인명)",
        "공급받는자-사업장 주소",
        "공급받는자-전화번호",
        "공급자-성명(대표자)",
        "공급자-상호(법인명)",
        "공급자-사업장 주소",
        "공급자-계좌번호",
        "공급자-등록번호",
        "YOUR REF NO",
        "합계금액",
        "Description",
        "VESSEL",
        "Amount",
        "U/Price",
        "출고일",
        "Unit",
        "Qty",
        "세액",
        "비고",
    ],
}

FIELD_LABEL_ALIASES = {
    "Grand Total": "Total Price",
    "Reference No.": "Our Ref No",
    "Our Ref. No.": "Our Ref No",
    "Your Ref. No.": "Your Ref No",
    "Vessel Name": "Vessel",
    "견적번호": "Our Ref No",
    "견적일자": "Date",
    "수신 회사": "To",
    "합계금액": "Total Price",
    "품목명": "Description",
    "수량": "Qty",
    "단가": "U/Price",
    "금액": "Amount",
    "delivery_time": "Terms & Conditions - Delivery time",
    "description": "Description",
    "total_price": "Total Price",
    "unit": "Unit",
    "unit_price": "U/Price",
    "quantity": "Qty",
}


def extract_predefined_document_fields(document_type: str | None, text: str) -> dict[str, Any]:
    """Best-effort extraction for known form types before optional LLM enrichment."""
    canonical_type = canonical_document_category(document_type)
    if canonical_type not in STRUCTURED_DOCUMENT_TYPES or not text.strip():
        return {}

    fields: dict[str, Any] = {}
    labels = sorted(FORM_FIELD_LABELS.get(canonical_type, []), key=len, reverse=True)
    for line in _form_candidate_lines(text, labels):
        for label in labels:
            value = _value_after_label(line, label)
            if not value:
                continue
            canonical_label = FIELD_LABEL_ALIASES.get(label, label)
            fields.setdefault(canonical_label, value)
            break

    if canonical_type == "quote":
        fields.update(_extract_quote_supplemental_fields(text, fields))

    line_items = _extract_line_items(canonical_type, text)
    if line_items:
        fields["line_items"] = line_items
    return filter_predefined_document_fields(canonical_type, fields)


def validated_predefined_document_fields(
    document_type: str | None,
    text: str,
    fields: dict[str, Any],
) -> tuple[dict[str, Any], list[str]]:
    """Filter, re-extract, and validate fixed-schema business document fields.

    The first pass keeps only the document schema. The second pass re-runs deterministic
    extraction against parsed text to recover fields the LLM/parser missed. The final
    pass records missing quote fields without adding non-schema substitutes.
    """
    canonical_type = canonical_document_category(document_type)
    filtered = filter_predefined_document_fields(canonical_type, fields)
    if text.strip() and canonical_type in PREDEFINED_DOCUMENT_FIELD_ORDER:
        recovered = extract_predefined_document_fields(canonical_type, text)
        if recovered:
            merged = dict(filtered)
            for key, value in recovered.items():
                if key not in merged or merged[key] in (None, "", [], {}):
                    merged[key] = value
            filtered = filter_predefined_document_fields(canonical_type, merged)
        if canonical_type == "quote":
            filtered = filter_predefined_document_fields(
                canonical_type,
                _recover_incomplete_quote_line_items(text, filtered, recovered),
            )
        if canonical_type == "rfq":
            filtered = filter_predefined_document_fields(
                canonical_type,
                _recover_incomplete_rfq_line_items(text, filtered, recovered),
            )
    warnings = _fixed_schema_validation_warnings(canonical_type, filtered, text)
    return filtered, warnings


def filter_predefined_document_fields(document_type: str | None, fields: dict[str, Any]) -> dict[str, Any]:
    """Keep only the fixed schema for known form documents.

    Quotation and RFQ documents may also keep extracted table rows as line_items.
    Unknown document types keep their existing field payload for non-form fallback display.
    """
    canonical_type = canonical_document_category(document_type)
    if canonical_type not in PREDEFINED_DOCUMENT_FIELD_ORDER:
        return fields
    allowed = set(PREDEFINED_DOCUMENT_FIELD_ORDER[canonical_type])
    normalized: dict[str, Any] = {}
    line_item_fields = _line_item_fields_for_document_type(canonical_type)
    _collect_predefined_fields(normalized, fields, allowed=allowed, fallback_label="", line_item_fields=line_item_fields)

    line_items = normalized.pop("line_items", None)
    ordered = {
        label: normalized[label]
        for label in PREDEFINED_DOCUMENT_FIELD_ORDER[canonical_type]
        if normalized.get(label) not in (None, "", [], {})
    }
    if canonical_type in {"quote", "rfq"} and isinstance(line_items, list) and line_items:
        ordered["line_items"] = line_items
    return ordered


def _collect_predefined_fields(
    normalized: dict[str, Any],
    item: Any,
    *,
    allowed: set[str],
    fallback_label: str,
    line_item_fields: tuple[str, ...],
) -> None:
    if item in (None, "", [], {}):
        return
    if isinstance(item, list):
        if fallback_label == "line_items":
            line_items = [_normalize_line_item(child, line_item_fields=line_item_fields) for child in item]
            line_items = [child for child in line_items if child]
            if line_items:
                normalized["line_items"] = line_items
            return
        if fallback_label == "dates":
            scalar_values = [_value_text(child) for child in item if not isinstance(child, (dict, list))]
            scalar_values = [value for value in scalar_values if value]
            if scalar_values:
                _set_allowed_field(normalized, "Date", scalar_values[0], allowed)
                return
        if fallback_label == "reference_numbers":
            scalar_values = [_value_text(child) for child in item if not isinstance(child, (dict, list))]
            scalar_values = [value for value in scalar_values if value]
            if scalar_values:
                _set_allowed_field(normalized, "Our Ref No", scalar_values[0], allowed)
            if len(scalar_values) > 1:
                _set_allowed_field(normalized, "Your Ref No", scalar_values[1], allowed)
            if scalar_values:
                return
        for child in item:
            _collect_predefined_fields(
                normalized,
                child,
                allowed=allowed,
                fallback_label="",
                line_item_fields=line_item_fields,
            )
        return
    if isinstance(item, dict):
        label = str(item.get("field_name") or item.get("name") or "").strip()
        field_value = item.get("value")
        if label and field_value not in (None, "", [], {}):
            _set_allowed_field(normalized, label, field_value, allowed)
            return
        for child_key, child_value in item.items():
            if child_key in {"field_name", "name", "value"}:
                continue
            _collect_predefined_fields(
                normalized,
                child_value,
                allowed=allowed,
                fallback_label=str(child_key),
                line_item_fields=line_item_fields,
            )
        return
    if fallback_label:
        _set_allowed_field(normalized, fallback_label, item, allowed)


def _set_allowed_field(normalized: dict[str, Any], label: str, value: Any, allowed: set[str]) -> None:
    exact_label = str(label or "").strip()
    if exact_label in allowed and value not in (None, "", [], {}):
        normalized[exact_label] = value


def _line_item_fields_for_document_type(document_type: str) -> tuple[str, ...]:
    if document_type == "rfq":
        return RFQ_LINE_ITEM_FIELDS
    if document_type == "quote":
        return QUOTE_LINE_ITEM_FIELDS
    return ()


def _normalize_line_item(item: Any, *, line_item_fields: tuple[str, ...]) -> dict[str, Any]:
    if not isinstance(item, dict):
        return {}
    normalized: dict[str, Any] = {}
    allowed = set(line_item_fields)
    for key, value in item.items():
        label = str(key or "").strip()
        if label in allowed and value not in (None, "", [], {}):
            normalized[label] = _quote_line_item_value(label, value)
    return normalized


def _quote_line_item_value(label: str, value: Any) -> Any:
    if label not in {"U/Price", "Amount"}:
        return value
    text = _value_text(value)
    if not text:
        return value
    return _strip_korean_won_unit(text)


def _strip_korean_won_unit(value: str) -> str:
    return re.sub(r"(?<=\d)\s*원(?![A-Za-z가-힣])", "", value).strip()


def _value_text(value: Any) -> str:
    if value in (None, "", [], {}):
        return ""
    return str(value).strip()


def _business_lines(text: str) -> list[str]:
    return [" ".join(line.split()) for line in text.splitlines() if line.strip()]


def _form_candidate_lines(text: str, labels: list[str]) -> list[str]:
    lines = _business_lines(text)
    candidates = list(lines)
    label_lookup = {label.casefold(): label for label in labels}
    for index, line in enumerate(lines):
        label = label_lookup.get(line.casefold())
        if not label:
            continue
        if index + 2 < len(lines) and lines[index + 1] in {":", "："}:
            candidates.append(f"{label} {lines[index + 2]}")
        elif index + 1 < len(lines) and lines[index + 1].startswith((": ", "： ")):
            candidates.append(f"{label} {lines[index + 1].lstrip(':： ')}")
    return candidates


def _value_after_label(line: str, label: str) -> str:
    escaped = re.escape(label).replace(r"\ ", r"\s+")
    match = re.match(rf"^{escaped}\s*(?:[:：\-]|\s)\s*(?P<value>.+)$", line, flags=re.IGNORECASE)
    if not match:
        return ""
    value = match.group("value").strip(" :|-\t")
    if not value or value.casefold() == label.casefold():
        return ""
    return value[:300]


def _extract_line_items(document_type: str, text: str) -> list[dict[str, str]]:
    lines = _business_lines(text)
    items: list[dict[str, str]] = []
    line_item_fields = _line_item_fields_for_document_type(document_type)
    if not line_item_fields:
        return []
    for index, line in enumerate(lines):
        headers = _split_table_row(line)
        normalized_headers = [FIELD_LABEL_ALIASES.get(header, header) for header in headers]
        if not {"Description", "Qty"}.issubset(set(normalized_headers)):
            continue
        for data_line in lines[index + 1 : index + 12]:
            if _looks_like_table_footer(data_line):
                break
            values = _split_table_row(data_line)
            if len(values) < 2 or len(values) != len(headers):
                compact = _parse_compact_line_item(document_type, data_line, normalized_headers)
                if compact:
                    items.append(compact)
                continue
            item = {
                header: _quote_line_item_value(header, value)
                for header, value in zip(normalized_headers, values, strict=True)
                if header in set(line_item_fields) and value
            }
            if item:
                items.append(item)
        break
    if not items and document_type == "quote":
        items = _extract_vertical_quote_line_items(lines)
    if not items and document_type == "rfq":
        items = _extract_vertical_rfq_line_items(lines)
    return items


def _extract_quote_supplemental_fields(text: str, existing: dict[str, Any]) -> dict[str, str]:
    fields: dict[str, str] = {}
    if not existing.get("Total Price"):
        total_match = re.search(
            r"\bTotal\s+Price\b\s*(?:[:：\-]|\s+is|\s+for\s+the\s+proposal\s+is|\s)\s*"
            r"(?P<value>(?:USD|KRW|EUR|JPY)?\s*[\d,]+(?:\.\d+)?(?:\s*\([^)]+\))?)",
            text,
            flags=re.IGNORECASE,
        )
        if total_match:
            fields["Total Price"] = " ".join(total_match.group("value").split())
    if not fields.get("Total Price"):
        total_price = _extract_vertical_quote_total_price(_business_lines(text))
        if total_price:
            fields["Total Price"] = total_price
    return fields


def _recover_incomplete_quote_line_items(
    text: str,
    fields: dict[str, Any],
    recovered: dict[str, Any],
) -> dict[str, Any]:
    expected_count = _expected_quote_line_item_count(text)
    if expected_count <= 0:
        return fields
    current_items = fields.get("line_items")
    current_count = len(current_items) if isinstance(current_items, list) else 0
    if current_count >= expected_count:
        return fields
    recovered_items = recovered.get("line_items")
    if not isinstance(recovered_items, list) or len(recovered_items) <= current_count:
        recovered_items = _extract_line_items("quote", text)
    if isinstance(recovered_items, list) and len(recovered_items) > current_count:
        merged = dict(fields)
        merged["line_items"] = recovered_items
        return merged
    return fields


def _recover_incomplete_rfq_line_items(
    text: str,
    fields: dict[str, Any],
    recovered: dict[str, Any],
) -> dict[str, Any]:
    expected_count = _expected_rfq_line_item_count(text)
    if expected_count <= 0:
        return fields
    current_items = fields.get("line_items")
    current_count = len(current_items) if isinstance(current_items, list) else 0
    if current_count >= expected_count:
        return fields
    recovered_items = recovered.get("line_items")
    if not isinstance(recovered_items, list) or len(recovered_items) <= current_count:
        recovered_items = _extract_line_items("rfq", text)
    if isinstance(recovered_items, list) and len(recovered_items) > current_count:
        merged = dict(fields)
        merged["line_items"] = recovered_items
        return merged
    return fields


def _fixed_schema_validation_warnings(
    document_type: str,
    fields: dict[str, Any],
    text: str = "",
) -> list[str]:
    if document_type == "rfq":
        return _fixed_rfq_validation_warnings(fields, text)
    if document_type != "quote":
        return []
    missing = [label for label in QUOTE_HEADER_FIELDS if fields.get(label) in (None, "", [], {})]
    line_items = fields.get("line_items")
    item_count = len(line_items) if isinstance(line_items, list) else 0
    expected_item_count = _expected_quote_line_item_count(text)
    if not isinstance(line_items, list) or not line_items:
        missing.extend(QUOTE_LINE_ITEM_FIELDS)
    else:
        for label in QUOTE_LINE_ITEM_FIELDS:
            if not any(isinstance(item, dict) and item.get(label) not in (None, "", [], {}) for item in line_items):
                missing.append(label)
    warnings = []
    if missing:
        warnings.append("fixed_quote_schema_missing:" + ",".join(dict.fromkeys(missing)))
    if expected_item_count and item_count < expected_item_count:
        warnings.append(f"fixed_quote_line_items_missing:expected={expected_item_count},actual={item_count}")
    return warnings


def _fixed_rfq_validation_warnings(fields: dict[str, Any], text: str) -> list[str]:
    line_items = fields.get("line_items")
    item_count = len(line_items) if isinstance(line_items, list) else 0
    expected_item_count = _expected_rfq_line_item_count(text)
    missing: list[str] = []
    if expected_item_count and (not isinstance(line_items, list) or not line_items):
        missing.extend(RFQ_LINE_ITEM_FIELDS)
    elif isinstance(line_items, list) and line_items:
        for label in RFQ_LINE_ITEM_FIELDS:
            if not any(isinstance(item, dict) and item.get(label) not in (None, "", [], {}) for item in line_items):
                missing.append(label)
    warnings = []
    if missing:
        warnings.append("fixed_rfq_schema_missing:" + ",".join(dict.fromkeys(missing)))
    if expected_item_count and item_count < expected_item_count:
        warnings.append(f"fixed_rfq_line_items_missing:expected={expected_item_count},actual={item_count}")
    return warnings


def _split_table_row(line: str) -> list[str]:
    if "|" in line:
        return [part.strip() for part in re.split(r"\s*\|\s*", line) if part.strip()]
    if "\t" in line:
        return [part.strip() for part in line.split("\t") if part.strip()]
    numbered_header_match = re.search(
        r"\bNo\.?\b.*\bDescription\b.*\bQty\b.*\bUnit\b.*\bU/?Price\b.*\bAmount\b",
        line,
        re.IGNORECASE,
    )
    if numbered_header_match:
        return ["No", "Description", "Qty", "Unit", "U/Price", "Amount"]
    header_match = re.search(r"\bDescription\b.*\bQty\b.*\bUnit\b.*\bU/?Price\b.*\bAmount\b", line, re.IGNORECASE)
    if header_match:
        return ["Description", "Qty", "Unit", "U/Price", "Amount"]
    if re.search(r"품목명\s+규격\s+수량\s+단가\s+금액", line):
        return ["Description", "Qty", "Unit", "U/Price", "Amount"]
    numbered_rfq_header_match = re.search(
        r"\bNo\.?\b.*\bDescription\b.*\bCode\b.*\bQty\b.*\bUnit\b",
        line,
        re.IGNORECASE,
    )
    if numbered_rfq_header_match:
        return ["No", "Description", "Code", "Qty", "Unit"]
    rfq_header_match = re.search(r"\bDescription\b.*\bCode\b.*\bQty\b.*\bUnit\b", line, re.IGNORECASE)
    if rfq_header_match:
        return ["Description", "Code", "Qty", "Unit"]
    normalized = re.sub(r"\s{2,}", "|", line.strip())
    parts = [part.strip() for part in normalized.split("|") if part.strip()]
    if len(parts) > 1:
        return parts
    return [line.strip()] if line.strip() else []


def _parse_compact_line_item(document_type: str, line: str, headers: list[str]) -> dict[str, str]:
    if document_type == "rfq":
        return _parse_compact_rfq_line_item(line, headers)
    return _parse_compact_quote_line_item(line, headers)


def _parse_compact_quote_line_item(line: str, headers: list[str]) -> dict[str, str]:
    if not set(QUOTE_LINE_ITEM_FIELDS).issubset(set(headers)):
        return {}
    match = re.match(
        r"(?:\d+\s+)?(?P<description>.+)\s+(?P<qty>\d+(?:\.\d+)?)\s+"
        r"(?P<unit>[A-Za-z가-힣]+)\s+(?P<uprice>(?:USD|KRW|EUR|JPY)?\s*[\d,]+(?:\.\d+)?(?:원)?)\s+"
        r"(?P<amount>(?:USD|KRW|EUR|JPY)?\s*[\d,]+(?:\.\d+)?(?:원)?)$",
        line.strip(),
        flags=re.IGNORECASE,
    )
    if not match:
        return {}
    return {
        "Description": " ".join(match.group("description").split()),
        "Qty": match.group("qty"),
        "Unit": match.group("unit"),
        "U/Price": _strip_korean_won_unit(" ".join(match.group("uprice").split())),
        "Amount": _strip_korean_won_unit(" ".join(match.group("amount").split())),
    }


def _parse_compact_rfq_line_item(line: str, headers: list[str]) -> dict[str, str]:
    header_set = set(headers)
    if not {"Description", "Code", "Qty", "Unit"}.issubset(header_set):
        return {}
    has_no = "No" in header_set
    pattern = (
        r"(?:(?P<no>\d+)\s+)?"
        r"(?P<description>.+?)\s+"
        r"(?P<code>[A-Za-z0-9][A-Za-z0-9._:/#\-]*)\s+"
        r"(?P<qty>\d+(?:\.\d+)?)\s+"
        r"(?P<unit>[A-Za-z가-힣]+)$"
    )
    match = re.match(pattern, line.strip())
    if not match:
        return {}
    item = {
        "Description": " ".join(match.group("description").split()),
        "Code": match.group("code"),
        "Qty": match.group("qty"),
        "Unit": match.group("unit"),
    }
    if has_no and match.group("no"):
        item = {"No": match.group("no"), **item}
    return item


def _looks_like_table_footer(line: str) -> bool:
    return bool(
        re.search(r"\b(Total\s+Price|Grand\s+Total|Terms\s*&\s*Conditions|Remark)\b", line, re.IGNORECASE)
        or re.search(r"(공급가액|부가세|합계금액|비고)", line)
        or re.match(r"PAGE\s+\d+\s*/", line, re.IGNORECASE)
    )


def _expected_quote_line_item_count(text: str) -> int:
    numbers = _quote_horizontal_line_item_numbers(_business_lines(text))
    numbers.extend(_quote_vertical_line_item_numbers(_business_lines(text)))
    return max(numbers) if numbers else 0


def _expected_rfq_line_item_count(text: str) -> int:
    numbers = _rfq_horizontal_line_item_numbers(_business_lines(text))
    numbers.extend(_rfq_vertical_line_item_numbers(_business_lines(text)))
    return max(numbers) if numbers else 0


def _quote_horizontal_line_item_numbers(lines: list[str]) -> list[int]:
    numbers: list[int] = []
    for index, line in enumerate(lines):
        headers = _split_table_row(line)
        normalized_headers = [FIELD_LABEL_ALIASES.get(header, header) for header in headers]
        if not {"No", "Description", "Qty"}.issubset(set(normalized_headers)):
            continue
        for data_line in lines[index + 1 : index + 80]:
            if _looks_like_table_footer(data_line):
                break
            values = _split_table_row(data_line)
            if values and re.fullmatch(r"\d+", values[0]):
                numbers.append(int(values[0]))
                continue
            match = re.match(r"\s*(?P<number>\d+)\s+\S+.*\s+\d+(?:\.\d+)?\s+\S+\s+", data_line)
            if match:
                numbers.append(int(match.group("number")))
        break
    return numbers


def _rfq_horizontal_line_item_numbers(lines: list[str]) -> list[int]:
    numbers: list[int] = []
    for index, line in enumerate(lines):
        headers = _split_table_row(line)
        normalized_headers = [FIELD_LABEL_ALIASES.get(header, header) for header in headers]
        if not {"No", "Description", "Code", "Qty", "Unit"}.issubset(set(normalized_headers)):
            continue
        for data_line in lines[index + 1 : index + 120]:
            if _looks_like_table_footer(data_line):
                break
            values = _split_table_row(data_line)
            if values and re.fullmatch(r"\d+", values[0]):
                numbers.append(int(values[0]))
                continue
            match = re.match(r"\s*(?P<number>\d+)\s+\S+.*\s+\S+\s+\d+(?:\.\d+)?\s+\S+\s*$", data_line)
            if match:
                numbers.append(int(match.group("number")))
        break
    return numbers


def _quote_vertical_line_item_numbers(lines: list[str]) -> list[int]:
    header_end = -1
    for index, line in enumerate(lines):
        if line.casefold() != "no":
            continue
        window = [item.casefold() for item in lines[index : index + 18]]
        if {"no", "description", "qty", "unit", "u/price", "amount"}.issubset(set(window)):
            amount_indexes = [offset for offset, value in enumerate(window) if value == "amount"]
            header_end = index + (amount_indexes[-1] if amount_indexes else 0)
            break
    if header_end < 0:
        return []
    tokens: list[str] = []
    for line in lines[header_end + 1 : header_end + 80]:
        if _looks_like_table_footer(line):
            break
        tokens.append(line)
    tokens = _dedupe_adjacent(tokens)
    numbers: list[int] = []
    for index, token in enumerate(tokens):
        if not re.fullmatch(r"\d+", token):
            continue
        if index + 4 >= len(tokens):
            continue
        qty = tokens[index + 2]
        unit = tokens[index + 3]
        unit_price = tokens[index + 4]
        amount = tokens[index + 5] if index + 5 < len(tokens) else ""
        if not amount or _looks_like_vertical_quote_item_start(tokens, index + 5) or _looks_like_table_footer(amount):
            amount = unit_price
        if _is_amount(qty) and unit.isalpha() and _is_amount(unit_price) and _is_amount(amount):
            numbers.append(int(token))
    return numbers


def _rfq_vertical_line_item_numbers(lines: list[str]) -> list[int]:
    numbers: list[int] = []
    for item in _extract_vertical_rfq_line_items(lines):
        number = str(item.get("No") or "")
        if re.fullmatch(r"\d+", number):
            numbers.append(int(number))
    return numbers


def _looks_like_vertical_quote_item_start(tokens: list[str], index: int) -> bool:
    if index + 4 >= len(tokens):
        return False
    return (
        bool(re.fullmatch(r"\d+", tokens[index]))
        and bool(tokens[index + 1].strip())
        and _is_amount(tokens[index + 2])
        and tokens[index + 3].isalpha()
        and _is_amount(tokens[index + 4])
    )


def _extract_vertical_quote_total_price(lines: list[str]) -> str:
    for index, line in enumerate(lines):
        if not re.search(r"\btotal\s+price\b", line, re.IGNORECASE):
            continue
        window = _dedupe_adjacent(lines[index + 1 : index + 10])
        currency = next((token for token in window if _is_currency(token)), "")
        amount = next((token for token in window if _is_amount(token)), "")
        vat_note = next((token for token in window if "vat" in token.casefold()), "")
        if amount:
            value = f"{currency} {amount}".strip()
            if vat_note:
                value = f"{value} {vat_note}"
            return value
    return ""


def _extract_vertical_quote_line_items(lines: list[str]) -> list[dict[str, str]]:
    header_start = -1
    header_end = -1
    for index, line in enumerate(lines):
        if line.casefold() != "description":
            continue
        window = [item.casefold() for item in lines[index : index + 16]]
        if {"description", "qty", "unit", "u/price", "amount"}.issubset(set(window)):
            header_start = index
            amount_indexes = [offset for offset, value in enumerate(window) if value == "amount"]
            header_end = index + (amount_indexes[-1] if amount_indexes else 0)
            break
    if header_start < 0:
        return []

    tokens: list[str] = []
    for line in lines[header_end + 1 : header_end + 80]:
        if _looks_like_table_footer(line):
            break
        tokens.append(line)
    tokens = _dedupe_adjacent(tokens)
    items: list[dict[str, str]] = []
    for index, token in enumerate(tokens):
        if not re.fullmatch(r"\d+", token):
            continue
        if index + 4 >= len(tokens):
            continue
        description = tokens[index + 1]
        qty = tokens[index + 2]
        unit = tokens[index + 3]
        unit_price = tokens[index + 4]
        amount = tokens[index + 5] if index + 5 < len(tokens) else ""
        if not amount or _looks_like_vertical_quote_item_start(tokens, index + 5) or _looks_like_table_footer(amount):
            amount = unit_price
        if not (_is_amount(qty) and unit.isalpha() and _is_amount(unit_price) and _is_amount(amount)):
            continue
        currency = next((item for item in tokens[index + 6 : index + 10] if _is_currency(item)), "")
        items.append(
            {
                "Description": description,
                "Qty": qty,
                "Unit": unit,
                "U/Price": _strip_korean_won_unit(f"{currency} {unit_price}".strip()),
                "Amount": _strip_korean_won_unit(f"{currency} {amount}".strip()),
            }
        )
    return items


def _extract_vertical_rfq_line_items(lines: list[str]) -> list[dict[str, str]]:
    header_end = -1
    for index, line in enumerate(lines):
        if line.casefold() != "no":
            continue
        window = [item.casefold() for item in lines[index : index + 12]]
        if {"no", "description", "code", "qty", "unit"}.issubset(set(window)):
            unit_indexes = [offset for offset, value in enumerate(window) if value == "unit"]
            header_end = index + (unit_indexes[-1] if unit_indexes else 0)
            break
    if header_end < 0:
        return []

    tokens: list[str] = []
    for line in lines[header_end + 1 : header_end + 120]:
        if _looks_like_table_footer(line):
            break
        tokens.append(line)
    tokens = _dedupe_adjacent(tokens)
    segmented_items = _extract_segmented_vertical_rfq_line_items(tokens)
    if segmented_items:
        return segmented_items
    items: list[dict[str, str]] = []
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if re.fullmatch(r"\d+", token) and index + 4 < len(tokens):
            description = tokens[index + 1]
            code = tokens[index + 2]
            qty = tokens[index + 3]
            unit = tokens[index + 4]
            if _looks_like_rfq_item(description, code, qty, unit):
                items.append({"No": token, "Description": description, "Code": code, "Qty": qty, "Unit": unit})
                index += 5
                continue
        if index + 4 < len(tokens) and re.fullmatch(r"\d+", tokens[index + 1]):
            description = token
            no = tokens[index + 1]
            code = tokens[index + 2]
            qty = tokens[index + 3]
            unit = tokens[index + 4]
            if _looks_like_rfq_item(description, code, qty, unit):
                items.append({"No": no, "Description": description, "Code": code, "Qty": qty, "Unit": unit})
                index += 5
                continue
        index += 1
    return items


def _extract_segmented_vertical_rfq_line_items(tokens: list[str]) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    search_start = 0
    expected_no = 1
    while search_start < len(tokens):
        start = _find_next_token(tokens, str(expected_no), search_start)
        if start < 0:
            break
        item: dict[str, str] = {}
        item_end = len(tokens)
        next_candidates = _token_indexes(tokens, str(expected_no + 1), start + 1)
        if next_candidates:
            for candidate_end in next_candidates:
                candidate = _parse_segmented_rfq_item(str(expected_no), tokens[start + 1 : candidate_end])
                if candidate:
                    item = candidate
                    item_end = candidate_end
                    break
        else:
            item = _parse_segmented_rfq_item(str(expected_no), tokens[start + 1 :])
        if not item:
            search_start = start + 1
            continue
        items.append(item)
        search_start = item_end
        expected_no += 1
    return items


def _find_next_token(tokens: list[str], expected: str, start: int) -> int:
    for index in range(start, len(tokens)):
        if tokens[index] == expected:
            return index
    return -1


def _token_indexes(tokens: list[str], expected: str, start: int) -> list[int]:
    return [index for index in range(start, len(tokens)) if tokens[index] == expected]


def _parse_segmented_rfq_item(no: str, segment: list[str]) -> dict[str, str]:
    if len(segment) < 4:
        return {}
    unit_index = -1
    qty_index = -1
    for index in range(len(segment) - 1, 0, -1):
        if not re.fullmatch(r"[A-Za-z가-힣]+", segment[index].strip()):
            continue
        if _is_amount(segment[index - 1]):
            unit_index = index
            qty_index = index - 1
            break
    if unit_index < 0 or qty_index < 2:
        return {}

    code_start = qty_index - 1
    while code_start > 0 and _looks_like_rfq_code_continuation(segment[code_start - 1], segment[code_start]):
        code_start -= 1
    description_parts = segment[:code_start]
    code_parts = segment[code_start:qty_index]
    if not description_parts or not code_parts:
        return {}
    description = " ".join(description_parts)
    code = "".join(code_parts) if any(part.endswith("-") for part in code_parts[:-1]) else " ".join(code_parts)
    return {
        "No": no,
        "Description": description,
        "Code": code,
        "Qty": segment[qty_index],
        "Unit": segment[unit_index],
    }


def _looks_like_rfq_code_continuation(previous: str, current: str) -> bool:
    if previous.endswith(("-", "/", ".")):
        return True
    return bool(re.fullmatch(r"[A-Za-z0-9._:/#\-]+", previous) and re.fullmatch(r"[A-Za-z0-9._:/#\-]+", current))


def _looks_like_rfq_item(description: str, code: str, qty: str, unit: str) -> bool:
    return bool(
        description.strip()
        and code.strip()
        and _is_amount(qty)
        and bool(re.fullmatch(r"[A-Za-z가-힣]+", unit.strip()))
    )


def _dedupe_adjacent(values: list[str]) -> list[str]:
    deduped: list[str] = []
    for value in values:
        if deduped and value == deduped[-1]:
            continue
        deduped.append(value)
    return deduped


def _is_currency(value: str) -> bool:
    return value.strip().upper() in {"KRW", "USD", "EUR", "JPY"}


def _is_amount(value: str) -> bool:
    return bool(re.fullmatch(r"\d+(?:,\d{3})*(?:\.\d+)?|\d+(?:\.\d+)?", value.strip()))
