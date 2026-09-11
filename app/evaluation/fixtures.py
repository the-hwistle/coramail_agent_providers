from __future__ import annotations

from pathlib import Path
from typing import Any


class SyntheticFixtureGenerator:
    """Materializes parser-ready files for synthetic attachment metadata."""

    def __init__(self, root: Path):
        self.root = root

    def generate(self, dataset: dict[str, Any]) -> dict[str, Any]:
        self.root.mkdir(parents=True, exist_ok=True)
        generated = 0
        for attachment in dataset.get("attachments") or []:
            path = self.root / attachment["filename"]
            self._write(path, attachment["synthetic_text"])
            attachment["storage_uri"] = str(path.resolve())
            attachment["file_size"] = path.stat().st_size
            generated += 1
        return {"fixtures": generated, "root": str(self.root.resolve())}

    def _write(self, path: Path, text: str) -> None:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            self._pdf(path, text)
        elif suffix == ".xlsx":
            self._xlsx(path, text)
        elif suffix == ".docx":
            self._docx(path, text)
        elif suffix == ".png":
            self._png(path, text)
        else:
            path.write_text(text, encoding="utf-8")

    @staticmethod
    def _pdf(path: Path, text: str) -> None:
        import fitz

        document = fitz.open()
        page = document.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 545, 792), text, fontsize=10)
        document.save(path)
        document.close()

    @staticmethod
    def _xlsx(path: Path, text: str) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Mail Request"
        for row_number, line in enumerate(text.splitlines(), start=1):
            key, separator, value = line.partition(":")
            sheet.cell(row=row_number, column=1, value=key.strip())
            sheet.cell(row=row_number, column=2, value=value.strip() if separator else line)
        workbook.save(path)

    @staticmethod
    def _docx(path: Path, text: str) -> None:
        from docx import Document

        document = Document()
        document.add_heading("Synthetic Mail Attachment", level=1)
        for line in text.splitlines():
            document.add_paragraph(line)
        document.save(path)

    @staticmethod
    def _png(path: Path, text: str) -> None:
        from PIL import Image, ImageDraw

        image = Image.new("RGB", (1400, 900), "white")
        draw = ImageDraw.Draw(image)
        y = 30
        for line in text.splitlines():
            draw.text((30, y), line, fill="black")
            y += 42
        image.save(path, format="PNG")
